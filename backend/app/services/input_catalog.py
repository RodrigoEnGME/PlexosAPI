"""Read and validate the documented SDDP-to-PLEXOS input dictionary."""

from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

EXPECTED_HEADERS = (
    "Atributo",
    "Nome........",
    "Tipo_Dato_SDDP",
    "Unidad_SDDP",
    "Archivo_SDDP",
    "Collection (Membership)",
    "Propiedad_PLEXOS",
    "Objeto_Plexos",
    "Subclase_Plexos",
    "Unidad_Plexos",
    "Comentario",
)


@dataclass(frozen=True, slots=True)
class InputFieldSpec:
    attribute: str
    sddp_name: str
    sddp_type: str
    sddp_unit: str
    source_file: str
    membership: str
    plexos_property: str
    plexos_object: str
    plexos_subclass: str
    plexos_unit: str
    comment: str
    source_sheet: str = ""
    source_row: int = 0

    def as_dict(self) -> dict[str, str]:
        return asdict(self)


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _load_workbook(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment/configuration error
        raise RuntimeError("openpyxl is required to read the input dictionary") from exc
    return load_workbook(path, read_only=True, data_only=True)


def load_input_dictionary(path: str | Path) -> list[InputFieldSpec]:
    """Load meaningful rows from the documented workbook.

    Blank formatted rows are ignored, while header drift is rejected so a
    revised workbook cannot silently change the catalog contract.
    """
    workbook_path = Path(path)
    if not workbook_path.is_file():
        raise FileNotFoundError(workbook_path)
    workbook = _load_workbook(workbook_path)
    try:
        if not workbook.sheetnames:
            raise ValueError("input dictionary has no worksheets")
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = tuple(_text(value) for value in next(rows, ()))
        if header != EXPECTED_HEADERS:
            raise ValueError(f"unexpected input dictionary headers: {header!r}")
        entries: list[InputFieldSpec] = []
        for row_number, row in enumerate(rows, start=2):
            values = tuple(_text(value) for value in row[: len(EXPECTED_HEADERS)])
            if not any(values):
                continue
            if len(values) != len(EXPECTED_HEADERS):
                raise ValueError("input dictionary contains a row with an invalid width")
            if not values[0] and not values[4]:
                continue
            entries.append(InputFieldSpec(*values, source_sheet=sheet.title, source_row=row_number))
        return entries
    finally:
        workbook.close()


def validate_input_dictionary(entries: list[InputFieldSpec]) -> None:
    """Reject duplicate field definitions for one source family."""
    seen: set[tuple[str, str, str]] = set()
    for entry in entries:
        key = (entry.source_file.casefold(), entry.sddp_name.casefold(), entry.attribute.casefold())
        if key in seen:
            raise ValueError(f"duplicate input dictionary field: {key}")
        seen.add(key)


def group_input_dictionary(entries: list[InputFieldSpec]) -> dict[str, list[InputFieldSpec]]:
    """Group field definitions by normalized source family."""
    grouped: dict[str, list[InputFieldSpec]] = {}
    for entry in entries:
        if entry.source_file in {"", "-"}:
            continue
        grouped.setdefault(entry.source_file.casefold(), []).append(entry)
    return grouped


def catalog_as_dict(path: str | Path) -> dict[str, object]:
    entries = load_input_dictionary(path)
    validate_input_dictionary(entries)
    grouped = group_input_dictionary(entries)
    return {
        "source": Path(path).name,
        "sheet_count": 1,
        "field_count": len(entries),
        "source_file_count": len(grouped),
        "fields": [entry.as_dict() for entry in entries],
    }
